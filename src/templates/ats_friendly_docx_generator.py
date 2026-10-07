#!/usr/bin/env python3
"""Single-column ATS-safe .docx generator for TJ's resume.

Accepts an optional `content` dict from jd_tailor.py for JD-tailored output.
When content=None the hardcoded base profile is used (default behavior).
"""

import os
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

_HERE = os.path.dirname(os.path.abspath(__file__))

GOLD  = RGBColor(0xC4, 0x9A, 0x3C)
DARK  = RGBColor(0x1A, 0x1A, 0x1A)
GRAY  = RGBColor(0x6B, 0x6B, 0x6B)


def _output_path(content=None):
    if content and content.get("meta"):
        company = content["meta"].get("company", "")
        slug    = content["meta"].get("role_slug", "")
        parts   = [p for p in [company.replace(" ", "_"), slug.replace("-", "_")] if p]
        suffix  = ("_" + "_".join(parts)) if parts else ""
        fname   = f"TJ_Tamilmani_Jayaraman_Resume{suffix}.docx"
    else:
        fname = "TJ_Tamilmani_Jayaraman_Resume.docx"
    return os.path.join(_HERE, "..", "..", "output", fname)


# ── Helpers ───────────────────────────────────────────────────────────────────
def set_font(run, size, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color


def para(doc, text, size=10, bold=False, color=None, italic=False,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=2):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_font(run, size, bold=bold, color=color, italic=italic)
    return p


def section_header(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(1)
    run = p.add_run(title.upper())
    set_font(run, 9, bold=True, color=GRAY)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "CCCCCC")
    pBdr.append(bottom)
    pPr.append(pBdr)


def bullet(doc, text, size=9.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.left_indent = Inches(0.15)
    run = p.add_run(text)
    set_font(run, size, color=DARK)


def job_entry(doc, title, company, dates, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(title)
    set_font(run, 10, bold=True, color=DARK)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    run2 = p2.add_run(company)
    set_font(run2, 9, bold=True, color=GOLD)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(2)
    run3 = p3.add_run(dates)
    set_font(run3, 8, italic=True, color=GRAY)

    for bt in bullets:
        bullet(doc, bt)


def edu_entry(doc, degree, institution):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(degree)
    set_font(run, 9, bold=True, color=DARK)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(3)
    run2 = p2.add_run(institution)
    set_font(run2, 8, color=GOLD)


# ── Base profile data ─────────────────────────────────────────────────────────
BASE = dict(
    tagline = (
        "Digital Engineering  |  Cloud Infrastructure  |  "
        "DevSecOps  |  Platform Engineering  |  AI/ML"
    ),
    summary = [
        "Strategic Technology Leader with 20+ years of experience leading enterprise-scale digital transformations and cloud modernization.",
        "Expert in building high-performance engineering cultures and delivering robust cloud foundations from ground zero to global production.",
        "Proven track record aligning technical roadmaps with business objectives — specializing in Digital Transformation, Cloud Engineering, FinOps Governance, DevSecOps Automation, and 99.99% system resiliency for multi-billion-dollar platforms.",
    ],
    jobs = [
        ("Director – Digital Engineering", "Southern Glazers Wine & Spirits",
         "May 2021 – Present  ·  Dallas, TX", [
            "Architected and executed comprehensive IT and cloud strategies scaling high-volume B2B commerce platforms to 24/7 global reliability.",
            "Led cross-functional rollout of CDN, IAM, and observability platforms, standardizing the enterprise technology stack.",
            "Negotiated multi-year vendor contracts exceeding $5M/year; managed $3M annual talent and staffing portfolio.",
            "Pioneered AI-driven operational excellence integrating foundational models and agentic workflows (Retrieval-Augmented Generation (RAG), Model Context Protocol (MCP)), boosting sales team productivity.",
            "Orchestrated organizational pivot from reactive AMS to proactive Site Reliability Engineering (SRE) — driving 35% Mean Time to Recovery (MTTR) reduction and 99.99% availability for $4B e-commerce platform.",
            "Championed FinOps maturity and transparent cloud governance; provided M&A technical due diligence.",
            "Strategic advisor to C-suite, aligning technology roadmaps with enterprise investment planning and long-term business objectives.",
        ]),
        ("Principal Engineer & Platform Leader", "Neiman Marcus Group",
         "Oct 2018 – May 2021  ·  Dallas, TX", [
            "Directed SRE and DevOps teams, standardizing methodologies and observability toolsets for global luxury retail platforms.",
            "Directed CDN transformation and mission-critical system upgrades to optimize the global digital footprint.",
            "Implemented rigorous capacity planning and rightsizing, driving significant cloud cost savings.",
            "Pioneered Robotic Process Automation (RPA) and Microservices automation (Java / Spring Boot) to eliminate toil and accelerate engineering velocity.",
            "Architected and scaled cloud-native AWS ecosystems (Serverless, Fargate, EC2) for high-availability enterprise infrastructure.",
        ]),
        ("Associate Director – Engineering", "Cognizant",
         "Apr 2010 – Aug 2018", [
            "Partnered with C-suite to scale digital commerce for Fortune 500 retail, driving enterprise adoption of Agile and DevOps.",
            "Managed $30M annual delivery portfolio; oversaw financial forecasts and KPIs across global programs.",
            "Built international Centers of Excellence (CoE) to standardize multi-regional delivery and foster cross-functional innovation.",
            "Led Waterfall to Agile (Scaled Agile Framework (SAFe) / Scrum / Kanban) and DevOps transformation; established TRAIN delivery framework.",
        ]),
        ("Manager – Engineering & Projects", "Atos Origin",
         "Oct 2009 – Apr 2010", [
            "Directed cross-continental teams delivering scalable automotive platforms for the Renault-Nissan Alliance.",
            "Spearheaded modular SOA architectures, significantly accelerating time-to-market for digital features.",
        ]),
        ("Director – Technology", "Mahindra Satyam",
         "Jan 2007 – Oct 2009", [
            "Managed multimillion-dollar project budgets for Ford and the World Bank, ensuring profitability and optimal resource allocation.",
            "Championed engineering excellence and mentorship programs, significantly increasing employee retention.",
        ]),
        ("Principal / Technology Lead", "Cognizant",
         "Sep 2002 – Jan 2007", [
            "Designed and implemented loan underwriting applications for AMEX (Java, J2EE, MQ, WebLogic).",
            "Designed customer service application for Chase Credit Card division (Visual Basic, Java, Oracle).",
        ]),
        ("Programmer", "Altosys Technologies",
         "Oct 2000 – Sep 2002", [
            "Developed client-server applications using Visual Basic, Oracle, MTS, and ASP.",
            "Designed application and database architecture, ER diagrams.",
        ]),
        ("Software Engineer", "Kaashyap Radiant Systems",
         "Mar 1999 – Sep 2000", [
            "Designed and built client-server applications using Visual Basic and Oracle.",
        ]),
    ],
    highlights = [
        ("Cloud Foundation at Scale",
         "Architected and deployed enterprise Cloud Foundation Platform in 12 weeks, enabling rapid multi-BU adoption."),
        ("99.99% Platform Availability",
         "Drove 35% MTTR reduction and 99.99% uptime for the $4B SGWS e-commerce platform via SRE transformation."),
        ("AI & GenAI Adoption",
         "Deployed Gen AI tools, RAG pipelines, and MCP agentic frameworks enterprise-wide, measurably boosting sales productivity."),
        ("FinOps Governance",
         "Established automated cloud governance workflows, driving fiscal accountability and material cloud cost reduction."),
        ("$30M Portfolio Leadership",
         "Managed $30M annual delivery portfolio at Cognizant — consistent on-time, on-budget delivery for Fortune 500 clients."),
        ("Vendor & Cost Optimization",
         "Negotiated multi-year contracts >$5M/year; strategic procurement achieving substantial toolchain cost savings."),
    ],
    competencies = (
        "Cloud Strategy & Governance  ·  Enterprise Digital Transformation  ·  "
        "DevSecOps & Automation  ·  FinOps & Cost Optimization  ·  SRE & Observability  ·  "
        "AI/ML & GenAI  ·  M&A Due Diligence  ·  Global Team Leadership  ·  Vendor Negotiation"
    ),
    technical_skills = [
        ("Cloud", "AWS (EC2, EKS, Fargate, Lambda, Bedrock, DynamoDB, API Gateway) · Azure (AKS, Cosmos DB, Functions)"),
        ("IaC & DevOps", "Terraform · CloudFormation · GitHub Actions · Kubernetes · Docker"),
        ("Languages", "Java · JavaScript · TypeScript · Python"),
        ("Data & AI", "GenAI · RAG · LangChain · LangGraph · Crew AI · Claude · Deep Learning · NLP"),
        ("Architecture", "Cloud-Native · Microservices · SOA · GraphQL · REST"),
    ],
    certifications = [
        "AWS Certified Solution Architect",
        "Project Management Professional (PMP)",
        "Certified Scrum Master",
        "Sun Certified Java Programmer",
        "Microsoft Certified Professional",
    ],
    education = [
        ("Post Graduate Program – AI & Machine Learning", "UT Austin  ·  2026"),
        ("Executive Program – Business Management", "IIM Calcutta, India  ·  2011"),
        ("Masters in Computer Applications", "Madurai Kamaraj University, India  ·  2001"),
        ("Bachelor of Science", "Madras University, India  ·  1997"),
    ],
)


def _resolve(content):
    if not content:
        return (
            BASE["tagline"], BASE["summary"], BASE["jobs"],
            BASE["highlights"], BASE["competencies"],
            BASE["technical_skills"], BASE["certifications"], BASE["education"],
        )

    tagline      = content.get("tagline", BASE["tagline"])
    summary      = content.get("summary", BASE["summary"])
    raw_exp      = content.get("experience", [])
    jobs         = [(e["title"], e["company"], e["dates"], e["bullets"]) for e in raw_exp] if raw_exp else BASE["jobs"]
    raw_hl       = content.get("highlights", [])
    highlights   = [(h["title"], h["desc"]) for h in raw_hl] if raw_hl else BASE["highlights"]
    competencies = content.get("competencies", BASE["competencies"])

    # technical_skills: AI returns bold-label strings; convert to (label, value) tuples for DOCX
    raw_skills = content.get("technical_skills", [])
    if raw_skills:
        skill_pairs = []
        for line in raw_skills:
            # Strip HTML bold tags if present
            line_clean = line.replace("<b>", "").replace("</b>", "")
            if ":" in line_clean:
                label, _, value = line_clean.partition(":")
                skill_pairs.append((label.strip(), value.strip()))
            else:
                skill_pairs.append(("", line_clean.strip()))
        technical_skills = skill_pairs
    else:
        technical_skills = BASE["technical_skills"]

    certifications = content.get("certifications", BASE["certifications"])
    raw_edu        = content.get("education", [])
    education      = [(e["degree"], e["institution"]) for e in raw_edu] if raw_edu else BASE["education"]

    return tagline, summary, jobs, highlights, competencies, technical_skills, certifications, education


# ── Build ─────────────────────────────────────────────────────────────────────
def build(content=None):
    out = _output_path(content)
    tagline, summary, jobs, highlights, competencies, technical_skills, certifications, education = _resolve(content)

    doc = Document()

    for section in doc.sections:
        section.top_margin    = Inches(0.55)
        section.bottom_margin = Inches(0.55)
        section.left_margin   = Inches(0.65)
        section.right_margin  = Inches(0.65)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"

    # ── Header ────────────────────────────────────────────────────────────────
    para(doc, "TAMILMANI JAYARAMAN (TJ)", size=22, bold=True, color=DARK,
         align=WD_ALIGN_PARAGRAPH.LEFT, space_after=2)
    para(doc, tagline, size=11, bold=True, color=GOLD, space_after=2)
    para(doc,
         "Dallas, TX  ·  tamizhmani2000@gmail.com  ·  (201) 290-9366  ·  "
         "https://www.linkedin.com/in/tamizh",
         size=9, color=GRAY, space_after=4)

    # ── Summary ───────────────────────────────────────────────────────────────
    section_header(doc, "Summary")
    for s in summary:
        bullet(doc, s)

    # ── Experience ────────────────────────────────────────────────────────────
    section_header(doc, "Experience")
    for title, company, dates, bullets_list in jobs:
        job_entry(doc, title, company, dates, bullets_list)

    # ── Professional Highlights ───────────────────────────────────────────────
    section_header(doc, "Professional Highlights")
    for title, desc in highlights:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(1)
        r1 = p.add_run(title + ": ")
        set_font(r1, 9, bold=True, color=DARK)
        r2 = p.add_run(desc)
        set_font(r2, 9, color=DARK)

    # ── Core Competencies ─────────────────────────────────────────────────────
    section_header(doc, "Core Competencies")
    para(doc, competencies, size=9, color=DARK, space_after=3)

    # ── Technical Skills ──────────────────────────────────────────────────────
    section_header(doc, "Technical Skills")
    for label, value in technical_skills:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        if label:
            r1 = p.add_run(label + ": ")
            set_font(r1, 9, bold=True, color=DARK)
        r2 = p.add_run(value)
        set_font(r2, 9, color=DARK)

    # ── Certifications ────────────────────────────────────────────────────────
    section_header(doc, "Certifications")
    for cert in certifications:
        bullet(doc, cert)

    # ── Education ─────────────────────────────────────────────────────────────
    section_header(doc, "Education")
    for deg, inst in education:
        edu_entry(doc, deg, inst)

    doc.save(out)
    print(f"Docx saved: {out}")
    return out


if __name__ == "__main__":
    build()
